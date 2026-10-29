import os
import json
import uuid
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import pypdf
import tempfile
from .database import get_connection, init_db, is_cloud_env

def create_sample_pptx(file_path: str):
    """Generate a realistic 12-slide MBA Product Management presentation."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide

    slides_info = [
        {
            "title": "Product Strategy: From Idea to Impact",
            "subtitle": "Build the right product, for the right users, at the right time.\nMBA Curriculum - Product Management Masterclass",
            "bullets": [
                "Strategic alignment between business goals, customer needs, and technical feasibility.",
                "How leading technology companies create enduring competitive moats.",
                "Frameworks: Kano Model, Blue Ocean, and Product-Market Fit."
            ]
        },
        {
            "title": "Understanding the Strategic Product Hierarchy",
            "subtitle": "Connecting high-level vision to ground-level sprint execution",
            "bullets": [
                "Company Vision: The ultimate destination and lasting purpose (10+ years).",
                "Product Strategy: The actionable plan, target segments, and competitive differentiators.",
                "Product Roadmap: Strategic themes and sequence of prioritized milestones.",
                "Sprint Backlog: User stories and tasks delivered by engineering squads."
            ]
        },
        {
            "title": "Customer Discovery & Problem Validation",
            "subtitle": "Fall in love with the problem, not your solution",
            "bullets": [
                "Continuous discovery habits: Conducting weekly touchpoints with target customers.",
                "The Mom Test principle: Ask about past behavior rather than hypothetical opinions.",
                "Identifying the 'Job-To-Be-Done' (JTBD) through structured inquiry.",
                "Distinguishing between 'vitamin' features and 'painkiller' solutions."
            ]
        },
        {
            "title": "Market Segmentation, Targeting & Positioning (STP)",
            "subtitle": "Designing products for specific customer archetypes",
            "bullets": [
                "Segmentation: Demographic, Psychographic, Behavioral, and Needs-based.",
                "Targeting: Evaluating Total Addressable Market (TAM), SAM, and SOM.",
                "Positioning Statement: For [target user] who [problem], our product is [category] that [key benefit].",
                "Avoiding the trap of building a generalized product for 'everyone'."
            ]
        },
        {
            "title": "Competitive Differentiation & Value Proposition",
            "subtitle": "Why should users choose your product over incumbent alternatives?",
            "bullets": [
                "Hamilton Helmer's 7 Powers of competitive advantage.",
                "Network effects: Direct vs Indirect network effects in marketplace platforms.",
                "High switching costs and deep proprietary data moats.",
                "Counter-positioning against established legacy incumbents."
            ]
        },
        {
            "title": "Prioritization Frameworks: RICE & Kano Models",
            "subtitle": "Making disciplined trade-offs with constrained engineering resources",
            "bullets": [
                "RICE Score = (Reach x Impact x Confidence) / Effort.",
                "Kano Model: Must-Haves (Basic needs), Performance attributes (Linear satisfaction), Delighters (Excitement).",
                "MoSCoW Prioritization: Must have, Should have, Could have, Won't have.",
                "Cost of Delay (CoD) and Weighted Shortest Job First (WSJF)."
            ]
        },
        {
            "title": "Defining Product-Market Fit (PMF) & Traction",
            "subtitle": "The Sean Ellis 40% test and retention curve dynamics",
            "bullets": [
                "Definition: When customer retention curves flatten out horizontally over time.",
                "The Sean Ellis Rule: If >40% of survey respondents would be 'very disappointed' without your product.",
                "Cohort Analysis: Tracking 30-day and 90-day retention across customer acquisition channels.",
                "Premature scaling: Spending heavily on paid marketing before validating retention."
            ]
        },
        {
            "title": "Unit Economics & Monetization Strategies",
            "subtitle": "Building a scalable, profitable business model",
            "bullets": [
                "Customer Lifetime Value (LTV) vs Customer Acquisition Cost (CAC).",
                "The 3:1 Rule: Healthy enterprise SaaS requires LTV / CAC >= 3.",
                "CAC Payback Period: Recovering acquisition investment within 12 months.",
                "Freemium vs Product-Led Growth (PLG) self-serve tiers vs Enterprise sales."
            ]
        },
        {
            "title": "Agile Delivery & Scrum Sprint Cycles",
            "subtitle": "Translating strategy into rapid, iterative customer value",
            "bullets": [
                "Sprint Planning, Daily Standups, Sprint Review, and Retrospectives.",
                "The Product Owner role: Defining acceptance criteria and user story definition.",
                "Managing technical debt alongside strategic customer-facing feature requests.",
                "Velocity tracking and release burndown metrics."
            ]
        },
        {
            "title": "Product Analytics & North Star Metric",
            "subtitle": "Guiding cross-functional teams with data-informed decisions",
            "bullets": [
                "The North Star Metric (NSM): The single metric that best captures core customer value.",
                "Examples: Airbnb (Nights Booked), Spotify (Time Spent Listening), Slack (Messages Sent in Teams).",
                "Input metrics: Actionable levers that influence the North Star Metric.",
                "A/B Testing: Establishing statistical significance and sample size calculation."
            ]
        },
        {
            "title": "Go-To-Market (GTM) Strategy & Product Launch",
            "subtitle": "Orchestrating marketing, sales, and customer success for launch",
            "bullets": [
                "Alpha and Beta customer cohorts: Gathering qualitative bug reports and usability feedback.",
                "Inbound vs Outbound distribution channels and organic search optimization.",
                "Sales enablement: Pitch decks, demo environments, battle cards against rivals.",
                "Customer onboarding: Minimizing Time-to-Value (TTV) to prevent early churn."
            ]
        },
        {
            "title": "Summary & Key Takeaways for MBA Exams",
            "subtitle": "Core review points for strategic product management",
            "bullets": [
                "Strategy is about focus: Choose what NOT to build with equal conviction.",
                "Never mistake outputs (shipped features) for outcomes (customer value & business growth).",
                "Maintain continuous empathy with end users while tracking unit economics.",
                "Review Questions: Describe the Kano Model and explain how a Delighter decays into a Must-Have."
            ]
        }
    ]

    for item in slides_info:
        slide = prs.slides.add_slide(blank_layout)
        
        # Title box
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.5), Inches(1.2))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = item["title"]
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
        
        # Subtitle
        if "subtitle" in item:
            p2 = tf.add_paragraph()
            p2.text = item["subtitle"]
            p2.font.size = Pt(14)
            p2.font.color.rgb = RGBColor(59, 130, 246) # Blue 500
            
        # Bullets box
        bodyBox = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.5), Inches(4.5))
        btf = bodyBox.text_frame
        btf.word_wrap = True
        for b in item["bullets"]:
            bp = btf.add_paragraph()
            bp.text = "• " + b
            bp.font.size = Pt(17)
            bp.space_after = Pt(14)
            bp.font.color.rgb = RGBColor(51, 65, 85) # Slate 700

    prs.save(file_path)

def seed_all_data():
    """Seed comprehensive MBA curriculum, subjects, files, quizzes, and notes."""
    init_db()
    conn = get_connection()
    cursor = conn.cursor()

    # Clear existing if any
    cursor.execute("DELETE FROM document_chunks_fts;")
    cursor.execute("DELETE FROM document_chunks;")
    cursor.execute("DELETE FROM quiz_attempts;")
    cursor.execute("DELETE FROM quizzes;")
    cursor.execute("DELETE FROM notes;")
    cursor.execute("DELETE FROM documents;")
    cursor.execute("DELETE FROM topics;")
    cursor.execute("DELETE FROM subjects;")

    # 1. Subjects
    subjects_data = [
        {"id": "sub_pm", "name": "Product Management", "semester": "Semester 1", "color": "#3B82F6", "icon": "Layers", "description": "Product discovery, strategy, UX design, sprint delivery and roadmaps."},
        {"id": "sub_mm", "name": "Marketing Management", "semester": "Semester 1", "color": "#EF4444", "icon": "Megaphone", "description": "Brand positioning, customer segmentation, digital marketing, and STP."},
        {"id": "sub_ba", "name": "Business Analytics", "semester": "Semester 1", "color": "#10B981", "icon": "BarChart2", "description": "Quantitative decision making, statistical modeling, and data-driven strategy."},
        {"id": "sub_fin", "name": "Finance", "semester": "Semester 2", "color": "#8B5CF6", "icon": "DollarSign", "description": "Corporate valuation, financial statements, DCF models, and capital budgeting."},
        {"id": "sub_om", "name": "Operations Management", "semester": "Semester 2", "color": "#F59E0B", "icon": "Truck", "description": "Supply chain optimization, inventory management, quality systems, and logistics."},
        {"id": "sub_ob", "name": "Organizational Behaviour", "semester": "Semester 3", "color": "#EC4899", "icon": "Users", "description": "Leadership dynamics, high-performing team culture, and organizational design."}
    ]

    for s in subjects_data:
        cursor.execute("""
        INSERT INTO subjects (id, name, semester, color, icon, description)
        VALUES (?, ?, ?, ?, ?, ?)
        """, (s["id"], s["name"], s["semester"], s["color"], s["icon"], s["description"]))

    # 2. Topics
    topics_data = [
        # PM Topics
        {"id": "top_ps", "subject_id": "sub_pm", "name": "Product Strategy", "status": "In Progress", "progress": 75, "order_index": 1},
        {"id": "top_pd", "subject_id": "sub_pm", "name": "Product Discovery", "status": "Understood", "progress": 85, "order_index": 2},
        {"id": "top_as", "subject_id": "sub_pm", "name": "Agile and Scrum", "status": "Completed", "progress": 100, "order_index": 3},
        {"id": "top_pr", "subject_id": "sub_pm", "name": "Product Roadmaps", "status": "Needs Revision", "progress": 40, "order_index": 4},
        
        # Marketing Topics
        {"id": "top_mm_mix", "subject_id": "sub_mm", "name": "Marketing Mix - Case Study", "status": "In Progress", "progress": 40, "order_index": 1},
        {"id": "top_mm_stp", "subject_id": "sub_mm", "name": "Market Segmentation & STP", "status": "Completed", "progress": 90, "order_index": 2},
        {"id": "top_mm_cb", "subject_id": "sub_mm", "name": "Consumer Behavior", "status": "Understood", "progress": 70, "order_index": 3},
        
        # Finance Topics
        {"id": "top_fin_fsa", "subject_id": "sub_fin", "name": "Financial Statement Analysis", "status": "In Progress", "progress": 25, "order_index": 1},
        {"id": "top_fin_cb", "subject_id": "sub_fin", "name": "Capital Budgeting", "status": "Needs Revision", "progress": 30, "order_index": 2},

        # Analytics Topics
        {"id": "top_ba_reg", "subject_id": "sub_ba", "name": "Regression Analysis", "status": "Completed", "progress": 100, "order_index": 1},
        {"id": "top_ba_pred", "subject_id": "sub_ba", "name": "Predictive Modeling", "status": "In Progress", "progress": 50, "order_index": 2},

        # Operations
        {"id": "top_om_sc", "subject_id": "sub_om", "name": "Supply Chain Logistics", "status": "Understood", "progress": 65, "order_index": 1},

        # OB
        {"id": "top_ob_lead", "subject_id": "sub_ob", "name": "Leadership Styles", "status": "Completed", "progress": 95, "order_index": 1}
    ]

    for t in topics_data:
        cursor.execute("""
        INSERT INTO topics (id, subject_id, name, status, progress, order_index)
        VALUES (?, ?, ?, ?, ?, ?)
        """, (t["id"], t["subject_id"], t["name"], t["status"], t["progress"], t["order_index"]))

    # 3. Create Sample PPTX file on disk
    if is_cloud_env():
        samples_dir = os.path.join(tempfile.gettempdir(), "sample_materials")
    else:
        samples_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), "data", "sample_materials")
    sample_pptx_path = os.path.join(samples_dir, "Product_Strategy.pptx")
    try:
        os.makedirs(samples_dir, exist_ok=True)
        if not os.path.exists(sample_pptx_path):
            create_sample_pptx(sample_pptx_path)
    except Exception as e:
        print(f"Notice generating sample pptx on cloud: {e}")

    # 4. Documents matching screenshots
    docs_data = [
        {
            "id": "doc_ps",
            "subject_id": "sub_pm",
            "topic_id": "top_ps",
            "name": "Product_Strategy.pptx",
            "original_filename": "Product_Strategy.pptx",
            "file_type": "PPTX",
            "file_size": 4200000,
            "file_path": sample_pptx_path,
            "semester": "Semester 1",
            "faculty": "Prof. R. Sharma",
            "description": "Comprehensive slide deck covering product strategy, market moats, PMF, and agile roadmapping.",
            "tags": json.dumps(["Important", "Exam Preparation"]),
            "slide_count": 12,
            "page_count": 12,
            "progress": 60
        },
        {
            "id": "doc_ms",
            "subject_id": "sub_mm",
            "topic_id": "top_mm_stp",
            "name": "Marketing_Segmentation.pdf",
            "original_filename": "Marketing_Segmentation.pdf",
            "file_type": "PDF",
            "file_size": 3800000,
            "file_path": sample_pptx_path, # re-use sample path for reference
            "semester": "Semester 1",
            "faculty": "Prof. A. Kapoor",
            "description": "Demographic, psychographic, and behavioral customer segmentation models and STP frameworks.",
            "tags": json.dumps(["Important", "Revision"]),
            "slide_count": 8,
            "page_count": 8,
            "progress": 40
        },
        {
            "id": "doc_fa",
            "subject_id": "sub_fin",
            "topic_id": "top_fin_fsa",
            "name": "Financial_Analysis.pdf",
            "original_filename": "Financial_Analysis.pdf",
            "file_type": "PDF",
            "file_size": 3800000,
            "file_path": sample_pptx_path,
            "semester": "Semester 2",
            "faculty": "Dr. V. Mehta",
            "description": "Three statement financial analysis, balance sheet ratios, cash flows, and working capital.",
            "tags": json.dumps(["Exam Preparation"]),
            "slide_count": 10,
            "page_count": 10,
            "progress": 25
        },
        {
            "id": "doc_as",
            "subject_id": "sub_pm",
            "topic_id": "top_as",
            "name": "Agile_Scrum.md",
            "original_filename": "Agile_Scrum.md",
            "file_type": "MD",
            "file_size": 1100000,
            "file_path": sample_pptx_path,
            "semester": "Semester 1",
            "faculty": "Prof. R. Sharma",
            "description": "Scrum ceremonies, product backlog grooming, story estimation, and sprint burndown metrics.",
            "tags": json.dumps(["Revision"]),
            "slide_count": 6,
            "page_count": 6,
            "progress": 80
        },
        {
            "id": "doc_on",
            "subject_id": "sub_om",
            "topic_id": "top_om_sc",
            "name": "Operations_Notes.docx",
            "original_filename": "Operations_Notes.docx",
            "file_type": "DOCX",
            "file_size": 1200000,
            "file_path": sample_pptx_path,
            "semester": "Semester 2",
            "faculty": "Dr. K. Iyer",
            "description": "Supply chain optimization, EOQ calculations, bottleneck theory, and lean manufacturing.",
            "tags": json.dumps(["Assignment"]),
            "slide_count": 5,
            "page_count": 5,
            "progress": 50
        },
        {
            "id": "doc_cs",
            "subject_id": "sub_ba",
            "topic_id": "top_ba_reg",
            "name": "Case_Study.txt",
            "original_filename": "Case_Study.txt",
            "file_type": "TXT",
            "file_size": 500000,
            "file_path": sample_pptx_path,
            "semester": "Semester 1",
            "faculty": "Prof. S. Verma",
            "description": "E-commerce customer churn predictive modeling dataset and case analysis.",
            "tags": json.dumps(["Important"]),
            "slide_count": 4,
            "page_count": 4,
            "progress": 70
        },
        {
            "id": "doc_rp",
            "subject_id": "sub_mm",
            "topic_id": "top_mm_cb",
            "name": "Research_Paper.pdf",
            "original_filename": "Research_Paper.pdf",
            "file_type": "PDF",
            "file_size": 2800000,
            "file_path": sample_pptx_path,
            "semester": "Semester 1",
            "faculty": "Prof. A. Kapoor",
            "description": "Behavioral economics in retail pricing and loss aversion heuristics.",
            "tags": json.dumps(["Revision"]),
            "slide_count": 7,
            "page_count": 7,
            "progress": 30
        }
    ]

    for d in docs_data:
        cursor.execute("""
        INSERT INTO documents (id, subject_id, topic_id, name, original_filename, file_type, file_size, file_path, semester, faculty, description, tags, slide_count, page_count, progress)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (d["id"], d["subject_id"], d["topic_id"], d["name"], d["original_filename"], d["file_type"], d["file_size"], d["file_path"], d["semester"], d["faculty"], d["description"], d["tags"], d["slide_count"], d["page_count"], d["progress"]))

    # 5. Extract and Index the PPTX file so RAG and Slide Viewer are immediately live!
    from .extractor import extract_pptx
    chunks = extract_pptx(sample_pptx_path)
    for c in chunks:
        chunk_id = str(uuid.uuid4())
        cursor.execute("""
        INSERT INTO document_chunks (id, document_id, chunk_index, slide_number, page_number, title, content, takeaways)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, (chunk_id, "doc_ps", c["slide_number"], c["slide_number"], c["page_number"], c["title"], c["content"], json.dumps(c["takeaways"])))
        
        try:
            cursor.execute("""
            INSERT INTO document_chunks_fts (chunk_id, document_id, title, content)
            VALUES (?, ?, ?, ?)
            """, (chunk_id, "doc_ps", c["title"], c["content"]))
        except Exception as e:
            pass

    # 6. Notes matching screenshot
    notes_data = [
        {"id": "note_1", "subject_id": "sub_pm", "topic_id": "top_ps", "title": "Product Lifecycle - Key Points", "note_type": "quick_revision", "content": "- Introduction: Heavy R&D investment, slow sales, negative profit.\n- Growth: Rapid market adoption, expanding margins, rising competition.\n- Maturity: Peak sales, price competition, brand loyalty defense.\n- Decline: Diminishing demand, harvest or divest strategy."},
        {"id": "note_2", "subject_id": "sub_mm", "topic_id": "top_mm_mix", "title": "Marketing Mix - 4Ps", "note_type": "quick_revision", "content": "1. Product: Features, design, brand name, packaging.\n2. Price: Cost-plus, value-based, penetration, skimming.\n3. Place: Direct sales, wholesale, omnichannel retail.\n4. Promotion: Advertising, PR, personal selling, digital performance."},
        {"id": "note_3", "subject_id": "sub_fin", "topic_id": "top_fin_fsa", "title": "Financial Ratios Summary", "note_type": "formula_sheet", "content": "Current Ratio = Current Assets / Current Liabilities\nDebt-to-Equity = Total Debt / Total Equity\nROE = Net Income / Shareholder Equity\nEBITDA Margin = EBITDA / Total Revenue"},
        {"id": "note_4", "subject_id": "sub_pm", "topic_id": "top_as", "title": "Agile vs Waterfall", "note_type": "quick_revision", "content": "- Agile: Iterative, flexible, customer feedback loops every 2 weeks.\n- Waterfall: Sequential, rigid stage gates, high risk of late failure."},
        {"id": "note_5", "subject_id": "sub_ba", "topic_id": "top_ba_reg", "title": "Porter's Five Forces", "note_type": "quick_revision", "content": "1. Threat of New Entrants\n2. Bargaining Power of Buyers\n3. Bargaining Power of Suppliers\n4. Threat of Substitute Products\n5. Competitive Rivalry within Industry"}
    ]
    for n in notes_data:
        cursor.execute("""
        INSERT INTO notes (id, subject_id, topic_id, title, note_type, content)
        VALUES (?, ?, ?, ?, ?, ?)
        """, (n["id"], n["subject_id"], n["topic_id"], n["title"], n["note_type"], n["content"]))

    # 7. Quizzes and Attempts
    from .ai_tutor import generate_quiz_for_topic
    sample_quiz = generate_quiz_for_topic(subject_id="sub_pm", topic_id="top_ps", document_id="doc_ps", difficulty="MBA Examination Level", num_questions=10)
    quiz_id = "quiz_pm_1"
    cursor.execute("""
    INSERT INTO quizzes (id, subject_id, topic_id, document_id, title, difficulty, format, questions_json)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, (quiz_id, "sub_pm", "top_ps", "doc_ps", "Product Strategy Comprehensive Quiz", "MBA Examination Level", "Multiple Choice", json.dumps(sample_quiz["questions"])))

    # Initial attempt record (score 8/10 -> 80% accuracy)
    cursor.execute("""
    INSERT INTO quiz_attempts (id, quiz_id, score, total_questions, correct_count, incorrect_count, accuracy, answers_json)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, ("att_1", quiz_id, 8, 10, 8, 2, 80, json.dumps([{"question_id": 1, "is_correct": True}, {"question_id": 2, "is_correct": True}])))

    conn.commit()
    conn.close()
    print("Sample MBA data seeded successfully!")

if __name__ == "__main__":
    seed_all_data()
